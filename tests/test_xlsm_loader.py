"""
Covers xlsm_loader.py's _normalize_cat_number() and the bug #29 fix
(_sheet_to_rows()'s blank-row early-exit + iter_rows() bulk reading).
"""
import time

import openpyxl

import xlsm_loader as xl


class TestNormalizeCatNumber:
    def test_none_passthrough(self):
        assert xl._normalize_cat_number(None) is None

    def test_float_from_excel_autotyping(self):
        # Excel auto-types a category-number cell with no leading zero
        # as a float, not text -- 10.0, 11.0, etc. Every resolver looks
        # categories up by a zero-padded 2-digit string.
        assert xl._normalize_cat_number(10.0) == "10"
        assert xl._normalize_cat_number(4.0) == "04"

    def test_string_with_stray_whitespace(self):
        assert xl._normalize_cat_number("04 ") == "04"

    def test_plain_digit_string_zero_padded(self):
        assert xl._normalize_cat_number("4") == "04"


def _make_phantom_range_sheet(real_rows=3, phantom_row=2000):
    """
    Reproduces the real bug #29 scenario in memory: a sheet whose
    reported max_row is dramatically larger than where real data ends --
    confirmed on a real file to happen from formatting/fill applied far
    beyond the actual content, not literal data. Touching a cell then
    clearing it reproduces the same shape (Excel registers the row as
    part of the used range even though the cell itself is blank).
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Fare Class"
    ws["B1"] = "OW/RT"
    for r in range(2, 2 + real_rows):
        ws.cell(row=r, column=1).value = f"FARE{r}"
        ws.cell(row=r, column=2).value = "2"
    ws.cell(row=phantom_row, column=1).value = "temp"
    ws.cell(row=phantom_row, column=1).value = None
    assert ws.max_row == phantom_row  # confirms the phantom range actually exists
    return wb, ws


class TestSheetToRowsBlankRowEarlyExit:
    def test_reads_only_real_rows_not_the_phantom_range(self):
        wb, ws = _make_phantom_range_sheet(real_rows=3, phantom_row=2000)
        rows = xl._sheet_to_rows(wb, ws.title, ["Fare Class", "OW/RT"])
        assert len(rows) == 3
        assert rows[0] == {"Fare Class": "FARE2", "OW/RT": "2"}

    def test_completes_quickly_despite_the_huge_reported_range(self):
        # Not a precise benchmark -- a loose regression guard. The old
        # per-cell ws.cell() loop took 5+ minutes (and was still not
        # finished) on a real 25,000-row version of this exact shape;
        # the fix should make this near-instant regardless of scale.
        wb, ws = _make_phantom_range_sheet(real_rows=3, phantom_row=20000)
        start = time.time()
        rows = xl._sheet_to_rows(wb, ws.title, ["Fare Class", "OW/RT"])
        elapsed = time.time() - start
        assert len(rows) == 3
        assert elapsed < 5.0, f"took {elapsed:.1f}s -- the early-exit likely regressed"

    def test_a_real_gap_inside_data_smaller_than_the_threshold_is_not_mistaken_for_the_end(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Fare Class"
        ws["B1"] = "OW/RT"
        ws.cell(row=2, column=1).value = "FIRST"
        ws.cell(row=2, column=2).value = "2"
        # A real gap of 10 blank rows (well under MAX_CONSECUTIVE_BLANK_ROWS) -- confirmed
        # every real file's data seen so far is contiguous, but the function
        # shouldn't treat a small gap as "end of data" even so.
        ws.cell(row=13, column=1).value = "AFTER_GAP"
        ws.cell(row=13, column=2).value = "3"
        rows = xl._sheet_to_rows(wb, ws.title, ["Fare Class", "OW/RT"])
        assert rows == [
            {"Fare Class": "FIRST", "OW/RT": "2"},
            {"Fare Class": "AFTER_GAP", "OW/RT": "3"},
        ]
