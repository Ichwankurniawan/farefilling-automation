"""
Reads category field labels DIRECTLY from the real template file, so the
template itself becomes the single source of truth for two things at once:
  1. The output schema (which fields exist, in what order) -- already used
     by template_writer.py's COLS dicts.
  2. What the AI is told to extract from free text -- instead of hand-
     written field descriptions in ai_specs/*.yaml (which can drift from
     the template), the AI now sees the EXACT labels a human filer would
     see in the spreadsheet (e.g. "Passenger Type", "MIN Age", "Travel").

This doesn't replace ai_specs/*.yaml entirely -- those still hold the
short category-level business instruction (what the category means).
But the field list + labels shown to the AI are now generated from the
template, not duplicated by hand.
"""
import openpyxl

SHEET_NAME = "(FINAL TEMPLATE) CAT 1-CAT 33 "

# List of header rows per category, top-to-bottom. Most categories have
# 1-2 header rows; CAT08 has 3 (group / sub-group / Charge1-Charge2
# sub-columns). For each template column, the LAST (bottom-most, most
# specific) non-empty label across these rows wins.
CATEGORY_HEADER_ROWS = {
    "CAT01": [12],
    "CAT02": [19, 20],
    "CAT03": [27, 28],
    "CAT04": [35, 36],
    "CAT08": [79, 80, 81],
}

_cache = {}


def get_field_labels(template_path, category, cols_map):
    """
    cols_map: the category's {internal_key: column_letter} dict (the same
    one template_writer.py uses to write cells) -- reused here so the
    column positions are defined in exactly one place.

    Returns: {internal_key: template_label} for every key in cols_map that
    has a non-empty label in the template. Returns {} (not an error) for
    a category not yet registered in CATEGORY_HEADER_ROWS -- field labels
    are an enrichment, not a hard requirement, so a resolver for a new
    category should still work without them.
    """
    if category not in CATEGORY_HEADER_ROWS:
        return {}

    cache_key = (template_path, category)
    if cache_key not in _cache:
        _cache[cache_key] = _read_header_rows(template_path, category)
    header_by_column = _cache[cache_key]

    labels = {}
    for key, col_letter in cols_map.items():
        label = header_by_column.get(col_letter)
        if label:
            labels[key] = label
    return labels


def _read_header_rows(template_path, category):
    wb = openpyxl.load_workbook(template_path)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb["(FINAL TEMPLATE) CAT 1-CAT 33 w"]
    rows = CATEGORY_HEADER_ROWS[category]

    result = {}
    for col in range(1, 91):
        col_letter = openpyxl.utils.get_column_letter(col)
        label = None
        for row in rows:  # bottom-most non-empty wins
            val = ws.cell(row=row, column=col).value
            if val:
                label = val
        if label:
            result[col_letter] = label
    return result
